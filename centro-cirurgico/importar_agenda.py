# -*- coding: utf-8 -*-
"""Importa a agenda diária do centro cirúrgico (PDF do hospital) para o Mapa_Cirurgico.

    python3 importar_agenda.py AGENDA_21.08.pdf Centro_Cirurgico_Escala.xlsx [saida.xlsx]

Sem argumentos, procura na pasta do próprio programa o PDF mais recente e a planilha mais
recente — é assim que os atalhos "Importar agenda" funcionam.

O que faz
  1. lê o PDF (pypdf, ou pdftotext se o pypdf não estiver instalado) e extrai sala, início,
     término, procedimento e cirurgião;
  2. traduz "Sala 1" para o código do posto usando a tabela POSTOS da aba Configuração;
  3. traduz o cirurgião para a especialidade usando a tabela CIRURGIÕES da aba Configuração,
     acrescentando ali, em branco, todo cirurgião ainda não cadastrado;
  4. substitui no Mapa_Cirurgico as linhas da data da agenda (as demais datas ficam intactas);
  5. aponta o Painel (Data_Mapa) para a data importada;
  6. grava uma NOVA versão do arquivo — o original nunca é sobrescrito.

O PDF traz dados de paciente (nome, prontuário, nascimento, convênio, leito). Nada disso é
lido nem gravado: só entram sala, horários, procedimento e cirurgião.
"""
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import date, datetime, time

from openpyxl import load_workbook

RE_HORA = re.compile(r"(\d{2}/\d{2}/\d{2})\s+(\d{2}:\d{2})\s*-\s*(\d{2}/\d{2}/\d{2})\s+(\d{2}:\d{2})")
RE_TIPO = re.compile(r"(?:Atendimento )?(?:Ambulatorial|Internado|Externo|Hospital Dia|"
                     r"Urgência|Emergência)")
CAB = ["Cirurgião", "Anestesista", "Cirurgia", "Tipo de Atendimento", "Observações", "Acomodação"]
MP_R1, MP_R2 = 2, 201          # faixa de dados do Mapa_Cirurgico
STATUS_PADRAO = "Agendada"


# ------------------------------------------------------------------ leitura do PDF
def _spans(linha):
    return [(m.start(), m.end(), m.group().strip()) for m in re.finditer(r"\S+(?: \S+)*", linha)]


def _colunas(hdr):
    """Posições das colunas a partir da linha de cabeçalho do relatório."""
    pos, busca = {}, 0
    for c in CAB:
        p = hdr.find(c, busca)
        if p >= 0:
            pos[c], busca = p, p + len(c)
    ordem = sorted(pos.items(), key=lambda kv: kv[1])
    return [(nome, ini, ordem[k + 1][1] if k + 1 < len(ordem) else 10 ** 6)
            for k, (nome, ini) in enumerate(ordem)]


def _distribui(linha, faixas):
    """Cada trecho da linha vai para a coluna com que tem maior sobreposição."""
    campos = {n: [] for n, _, _ in faixas}
    for a, b, txt in _spans(linha):
        melhor, best = None, 0
        for nome, ini, fim in faixas:
            ov = min(b, fim) - max(a, ini)
            if ov > best:
                melhor, best = nome, ov
        if melhor:
            campos[melhor].append(txt)
    return {n: " ".join(v) for n, v in campos.items()}


def texto_do_pdf(pdf):
    """Texto do PDF preservando as colunas. Usa o pypdf; sem ele, cai no pdftotext."""
    try:
        from pypdf import PdfReader
    except ImportError:
        if shutil.which("pdftotext") is None:
            sys.exit("Instale o leitor de PDF:  pip install pypdf")
        return subprocess.run(["pdftotext", "-layout", pdf, "-"],
                              capture_output=True, text=True).stdout
    leitor = PdfReader(pdf)
    if leitor.is_encrypted:          # o relatório do hospital vem protegido contra cópia
        leitor.decrypt("")
    return "\n".join(p.extract_text(extraction_mode="layout") for p in leitor.pages)


def ler_agenda(pdf):
    """Extrai as cirurgias do PDF do mapa cirúrgico."""
    return ler_texto(texto_do_pdf(pdf))


def ler_texto(texto):
    """Mesma extração a partir do texto já convertido (usada também nos testes)."""
    L = texto.split("\n")
    sala, out = None, []
    for i, ln in enumerate(L):
        s = ln.strip()
        if re.match(r"^Sala\s+\S+$", s, re.I):
            sala = s
            continue
        m = RE_HORA.search(ln)
        if not m:
            continue
        d1, h1, _, h2 = m.groups()
        hi = next((k for k in range(i + 1, min(i + 4, len(L)))
                   if "Cirurgião" in L[k] and "Cirurgia" in L[k]), None)
        reg = dict(sala=sala, data=datetime.strptime(d1, "%d/%m/%y").date(),
                   ini=h1, fim=h2, cirurgiao="", procedimento="", tipo="")
        if hi is not None:
            faixas = _colunas(L[hi])
            c = _distribui(L[hi + 1], faixas)
            reg.update(cirurgiao=c.get("Cirurgião", ""), procedimento=c.get("Cirurgia", ""),
                       tipo=c.get("Tipo de Atendimento", ""))
            for k in range(hi + 2, min(hi + 6, len(L))):   # nome do procedimento em 2 linhas
                nx = L[k]
                if not nx.strip() or RE_HORA.search(nx) or "Equipamento" in nx \
                        or "Início - Término" in nx:
                    break
                ex = _distribui(nx, faixas).get("Cirurgia", "")
                if ex and not ex.startswith("Item Sem"):
                    reg["procedimento"] += " " + ex
            reg["procedimento"] = re.sub(r"\s+", " ", reg["procedimento"]).strip()
            if not reg["tipo"]:      # vocabulário fechado: recupera o tipo de onde tiver caído
                for campo in ("procedimento",):
                    mt = RE_TIPO.search(reg[campo])
                    if mt:
                        reg["tipo"] = mt.group().strip()
                        reg[campo] = (reg[campo][:mt.start()] + " " + reg[campo][mt.end():]).strip()
        out.append(reg)
    return out


# ------------------------------------------------------------------ escrita na planilha
def _faixa(wb, nome):
    """Linhas (primeira, última) e coluna de um nome definido de coluna inteira."""
    dest = list(wb.defined_names[nome].destinations)[0]
    ws, ref = wb[dest[0]], dest[1].replace("$", "")
    a, b = (ref.split(":") + [ref])[:2]
    return ws, int(re.sub(r"\D", "", a)), int(re.sub(r"\D", "", b)), re.sub(r"\d", "", a)


def _hhmm(t):
    return time(int(t[:2]), int(t[3:]))


def importar(pdf, entrada, saida=None, cirs=None):
    cirs = cirs if cirs is not None else ler_agenda(pdf)
    if not cirs:
        sys.exit("Nenhuma cirurgia encontrada no PDF — confira se é o mapa do centro cirúrgico.")
    dia = cirs[0]["data"]
    wb = load_workbook(entrada)
    cfg, mapa = wb["Configuração"], wb["Mapa_Cirurgico"]

    ws, p1, p2, _ = _faixa(wb, "Postos_Cod")
    postos = {str(ws.cell(row=r, column=2).value).strip(): ws.cell(row=r, column=1).value
              for r in range(p1, p2 + 1) if ws.cell(row=r, column=1).value}
    ws, c1, c2, _ = _faixa(wb, "Cir_Nome")
    cir_esp = {str(ws.cell(row=r, column=1).value).strip(): (ws.cell(row=r, column=2).value or "")
               for r in range(c1, c2 + 1) if ws.cell(row=r, column=1).value}

    # linhas de outras datas são preservadas
    outras = []
    for r in range(MP_R1, MP_R2 + 1):
        v = [mapa.cell(row=r, column=c).value for c in range(1, 9)]
        d = v[0].date() if isinstance(v[0], datetime) else v[0]
        if v[1] and d != dia:
            outras.append(v)

    novas, sem_sala, sem_esp = [], set(), set()
    for c in cirs:
        cod = postos.get(c["sala"])
        if not cod:
            sem_sala.add(c["sala"])
        esp = cir_esp.get(c["cirurgiao"], "")
        if not esp:
            sem_esp.add(c["cirurgiao"])
        novas.append([dia, cod or c["sala"], _hhmm(c["ini"]), _hhmm(c["fim"]),
                      esp, c["procedimento"], c["cirurgiao"], STATUS_PADRAO])

    linhas = sorted(outras + novas, key=lambda v: (str(v[0]), str(v[1]), str(v[2])))
    if len(linhas) > MP_R2 - MP_R1 + 1:
        sys.exit("Mapa_Cirurgico tem %d linhas e a importação precisa de %d."
                 % (MP_R2 - MP_R1 + 1, len(linhas)))
    for i in range(MP_R2 - MP_R1 + 1):
        r = MP_R1 + i
        for j in range(8):
            mapa.cell(row=r, column=1 + j).value = linhas[i][j] if i < len(linhas) else None

    # cirurgiões novos entram na tabela, em branco, para a coordenação classificar
    livres = [r for r in range(c1, c2 + 1) if not ws.cell(row=r, column=1).value]
    for nome in sorted(n for n in sem_esp if n and n not in cir_esp):
        if not livres:
            print("! tabela de cirurgiões cheia: %s não foi cadastrado" % nome)
            break
        ws.cell(row=livres.pop(0), column=1).value = nome

    mapa.print_area = "A1:N%d" % max(MP_R1, MP_R1 + len(linhas) - 1)

    d_ws, d_r, _, d_c = _faixa(wb, "Data_Mapa") if "Data_Mapa" in wb.defined_names else (None,) * 4
    if d_ws is not None:
        d_ws["%s%d" % (d_c, d_r)] = dia

    if not saida:
        # tira um "_AAAA-MM-DD" e um "_vN" que já estejam no nome, para não empilhar sufixos
        base = re.sub(r"_\d{4}-\d{2}-\d{2}(_v\d+)?(?=\.xlsx$)", "", entrada)
        saida = base.replace(".xlsx", "_%s.xlsx" % dia.strftime("%Y-%m-%d"))
        v = 2   # nunca sobrescreve: reimportar o mesmo dia gera _v2, _v3...
        while os.path.exists(saida):
            saida = base.replace(".xlsx", "_%s_v%d.xlsx" % (dia.strftime("%Y-%m-%d"), v))
            v += 1
    if os.path.abspath(saida) == os.path.abspath(entrada):
        sys.exit("A saída não pode ser o arquivo de entrada.")
    wb.save(saida)

    turno_fim = cfg["B12"].value
    lim = turno_fim.hour * 60 + turno_fim.minute if isinstance(turno_fim, time) else 13 * 60
    fora = [c for c in cirs if int(c["fim"][:2]) * 60 + int(c["fim"][3:]) > lim]
    print("data da agenda ......: %s" % dia.strftime("%d/%m/%Y"))
    print("cirurgias importadas : %d em %d salas" % (len(novas), len({c['sala'] for c in cirs})))
    print("outras datas mantidas: %d linhas" % len(outras))
    if sem_sala:
        print("! sala sem cadastro .: %s" % ", ".join(sorted(sem_sala)))
    if sem_esp:
        print("! sem especialidade .: %s" % ", ".join(sorted(n for n in sem_esp if n)))
        print("  (cadastre a especialidade na aba Configuração, seção 8 — só uma vez por cirurgião)")
    if fora:
        print("! passam do plantão .: %d cirurgia(s) terminam depois de %02d:%02d"
              % (len(fora), lim // 60, lim % 60))
    print("salvo: %s" % saida)
    return saida


def _mais_recente(pasta, ext, ignorar=()):
    alvos = [os.path.join(pasta, f) for f in os.listdir(pasta)
             if f.lower().endswith(ext) and not f.startswith("~$")
             and not any(t in f for t in ignorar)]
    return max(alvos, key=os.path.getmtime) if alvos else None


def main(args):
    abrir = "--abrir" in args
    args = [a for a in args if a != "--abrir"]
    saida = _rodar(args)
    if abrir and saida and hasattr(os, "startfile"):   # Windows: abre a planilha pronta
        os.startfile(saida)                            # noqa: S606
    return saida


def _rodar(args):
    if len(args) >= 2:
        return importar(args[0], args[1], args[2] if len(args) > 2 else None)
    pasta = os.path.dirname(os.path.abspath(__file__))
    pdf = args[0] if args else _mais_recente(pasta, ".pdf")
    planilha = _mais_recente(pasta, ".xlsx")
    if not pdf:
        sys.exit("Nenhum PDF encontrado em %s. Copie a agenda para essa pasta." % pasta)
    if not planilha:
        sys.exit("Nenhuma planilha .xlsx encontrada em %s." % pasta)
    print("agenda ..............: %s" % os.path.basename(pdf))
    print("planilha ............: %s" % os.path.basename(planilha))
    return importar(pdf, planilha)


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except SystemExit:
        raise
    except Exception as e:                       # mensagem legível em vez de traceback
        sys.exit("Não deu para importar: %s" % e)
