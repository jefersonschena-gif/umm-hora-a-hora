# -*- coding: utf-8 -*-
"""Teste 8 — Painel: números, linha do tempo do plantão, fila de ações e os dois rankings,
recalculados por fora a partir do Mapa do Dia e da Agenda do Dia."""
import os
import sys
from datetime import datetime, time

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from abas import ABA_AGENDA, ABA_CFG, ABA_EQUIPE, ABA_MAPA, ABA_PAINEL

ARQ = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, "Centro_Cirurgico_Escala.xlsx")
wb = openpyxl.load_workbook(ARQ, data_only=True)
pai, mapa, age, cfg, eqp = (wb[ABA_PAINEL], wb[ABA_MAPA], wb[ABA_AGENDA], wb[ABA_CFG], wb[ABA_EQUIPE])

PAI_R1, N_POS, POS_R1 = 10, 16, 45
COL_NOME, COL_DUPLA, N_SLOT, SLOT = 11, 10, 36, 10
COL_T0 = COL_NOME + COL_DUPLA + 1
N_COL = COL_NOME + COL_DUPLA + N_SLOT
DIA_R1 = 12
DIA_R2 = DIA_R1 + N_POS - 1
ACO_P1 = DIA_R2 + 4
ACO_P2 = ACO_P1 + 7
ENC_P1 = ACO_P2 + 4
ATR_P1 = ENC_P1 + 6
KPI_INI = [1, 10, 17, 24, 33, 40, 47]
JANELA = cfg["B17"].value or 0
LIMPEZA = cfg["B16"].value or 0
T_INI = cfg["B11"].value
BASE = T_INI.hour * 60 + T_INI.minute if isinstance(T_INI, time) else 420

N = lambda v: v if isinstance(v, (int, float)) else 0   # célula vazia ou texto conta como 0

res, ok = [], 0
def chk(nome, esperado, obtido):
    global ok
    bate = esperado == obtido
    ok += bate
    res.append("%-46s esperado %-30s obtido %-30s %s"
               % (nome, esperado, obtido, "ok" if bate else "DIVERGE"))

# --- o que o Mapa do Dia diz, linha a linha
LIN = []
for i in range(N_POS):
    r = PAI_R1 + i
    nome = mapa.cell(row=r, column=1).value
    if not nome:
        continue
    LIN.append(dict(i=i, nome=nome, cod=mapa.cell(row=r, column=14).value,
                    pede=mapa.cell(row=r, column=2).value,
                    sit=mapa.cell(row=r, column=6).value,
                    alerta=mapa.cell(row=r, column=7).value,
                    cir=mapa.cell(row=r, column=8).value,
                    cabe=mapa.cell(row=r, column=11).value,
                    atraso=mapa.cell(row=r, column=12).value,
                    t1=mapa.cell(row=r, column=4).value,
                    t2=mapa.cell(row=r, column=5).value,
                    encaixe=next(cfg.cell(row=x, column=11).value for x in range(POS_R1, POS_R1 + N_POS)
                                 if cfg.cell(row=x, column=2).value == nome),
                    tipo=next(cfg.cell(row=x, column=3).value for x in range(POS_R1, POS_R1 + N_POS)
                              if cfg.cell(row=x, column=2).value == nome)))

# --- 1. números do dia
num = lambda k: pai.cell(row=9, column=KPI_INI[k]).value
chk("ambientes abertos", sum(1 for l in LIN if N(l["pede"]) > 0), num(1))
chk("prontos", sum(1 for l in LIN if l["sit"] == "OK"), num(2))
chk("com pendência", sum(1 for l in LIN if l["sit"] not in (None, "", "OK", "NÃO OPERA")), num(3))
chk("disponíveis", sum(1 for r in range(2, 32)
                       if eqp.cell(row=r, column=5).value == "Disponível"), num(4))
chk("maior atraso", max([N(l["atraso"]) for l in LIN] + [0]), num(6))

# --- 2. linha do tempo: recalculada da agenda do dia
DATA = pai["G4"].value
DATA = DATA.date() if isinstance(DATA, datetime) else DATA
CIRS = {}
for r in range(2, 202):
    sala, ini, fim = (age.cell(row=r, column=2).value, age.cell(row=r, column=3).value,
                      age.cell(row=r, column=4).value)
    d = age.cell(row=r, column=1).value
    d = d.date() if isinstance(d, datetime) else d
    if not sala or d != DATA or N(age.cell(row=r, column=10).value) <= 0:
        continue
    a = (ini.hour * 60 + ini.minute - BASE) % 1440
    CIRS.setdefault(sala, []).append((a, a + N(age.cell(row=r, column=9).value)))

for l in LIN:
    r = DIA_R1 + l["i"]
    chk("linha do tempo %s: nome" % l["nome"], l["nome"], pai.cell(row=r, column=1).value)
    chk("linha do tempo %s: situação" % l["nome"], l["sit"], pai.cell(row=r, column=N_COL + 1).value)
    esperado = []
    for t in range(N_SLOT):
        t0, tf = t * SLOT, (t + 1) * SLOT
        v = 0
        # posto de apoio não tem linha do tempo: a barra é só das salas cirúrgicas
        for a, b in (CIRS.get(l["cod"], []) if l["tipo"] == "Sala cirúrgica" else []):
            if a < tf and b > t0:
                v = 2
            elif v == 0 and b < tf and b + LIMPEZA > t0:
                v = 1
        esperado.append(v)
    obtido = [N(pai.cell(row=r, column=COL_T0 + t).value) for t in range(N_SLOT)]
    chk("linha do tempo %s: barras" % l["nome"], esperado, obtido)

# --- 3. fila de ações
GRAV = {"SEM EQUIPE": 900, "FALTA HABILIDADE": 800, "GENTE DEMAIS": 700, "INCOMPLETO": 650,
        "EXTRA": 600, "ATENÇÃO": 400, "SEM AVALIAÇÃO": 300}
chk("nenhum ambiente com pendência mostra alerta OK", [],
    [l["nome"] for l in LIN if l["sit"] in GRAV and l["alerta"] == "OK"])
matriz_vazia = sum(1 for r in range(2, 32) for c in range(8, 28)
                   if str(eqp.cell(row=r, column=c).value or "").strip().upper() == "X") == 0
esperadas = []
for l in LIN:
    g = GRAV.get(l["sit"], 0)
    if l["sit"] == "SEM EQUIPE" and N(l["cir"]) <= 0:
        g = 850
    if l["sit"] == "SEM AVALIAÇÃO" and matriz_vazia:
        g = 0
    if g:
        esperadas.append((g, "%s · %s" % (l["nome"], l["alerta"])))
notas = [pai.cell(row=r, column=N_COL + 1).value for r in range(ACO_P1, ACO_P2 + 1)
         if pai.cell(row=r, column=3).value]
mostradas = [pai.cell(row=r, column=3).value for r in range(ACO_P1, ACO_P2 + 1)
             if pai.cell(row=r, column=3).value]
globais = sum(1 for n in notas if n >= 970)
chk("fila em ordem decrescente de gravidade", True, notas == sorted(notas, reverse=True))
chk("ações de ambiente listadas", min(len(esperadas), 8 - globais), len(mostradas) - globais)
chk("primeira ação de ambiente confere",
    sorted(esperadas, reverse=True)[0][1] if esperadas else None,
    mostradas[globais] if len(mostradas) > globais else None)
maior = max(notas) if notas else 0
chk("faixa do dia", ("AÇÃO NECESSÁRIA" if maior >= 900 else "ATENÇÃO" if maior >= 600
                     else "QUASE PRONTO" if maior >= 1 else "DIA FECHADO — nada pendente"),
    pai["A6"].value)

# --- 4. onde encaixar e risco de atraso
cand = sorted([l for l in LIN if l["encaixe"] == "Sim" and isinstance(l["cabe"], (int, float))
               and l["cabe"] >= JANELA], key=lambda l: -l["cabe"])[:3]
for i in range(3):
    r = ENC_P1 + i
    chk("encaixe %d ambiente" % (i + 1), cand[i]["nome"] if i < len(cand) else None,
        pai.cell(row=r, column=3).value)
    chk("encaixe %d cabe até" % (i + 1), cand[i]["cabe"] if i < len(cand) else None,
        pai.cell(row=r, column=33).value)
chk("nenhuma sala que não aceita encaixe no ranking", [],
    [pai.cell(row=ENC_P1 + i, column=3).value for i in range(3)
     for l in LIN if l["encaixe"] != "Sim" and l["nome"] == pai.cell(row=ENC_P1 + i, column=3).value])

atr = sorted([l for l in LIN if N(l["atraso"]) > 0], key=lambda l: -l["atraso"])[:3]
for i in range(3):
    r = ATR_P1 + i
    chk("atraso %d ambiente" % (i + 1), atr[i]["nome"] if i < len(atr) else None,
        pai.cell(row=r, column=3).value)
    chk("atraso %d minutos" % (i + 1), atr[i]["atraso"] if i < len(atr) else None,
        pai.cell(row=r, column=15).value)

print("\n".join(res))
print("\nTESTE 8 (painel): %d/%d" % (ok, len(res)))
sys.exit(0 if ok == len(res) else 1)
