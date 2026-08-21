# -*- coding: utf-8 -*-
"""Teste 10 — matriz de afinidade da aba Equipe.

A matriz cruza os nomes: X quer dizer que as duas podem dividir o mesmo posto. O que este
teste cobra, do arquivo gerado até a distribuição:
  · a matriz nasce marcada, inclusive nas linhas ainda sem nome (quem for cadastrado depois
    já entra podendo trabalhar com todo mundo);
  · apagar o X de um lado só já proíbe a dupla — a planilha lê as duas células;
  · quem está proibida com todas as disponíveis fica sozinha no posto, não é escalada à força;
  · a fórmula do Mapa do Dia acusa DUPLA VETADA quando a dupla é formada à mão.
"""
import os, shutil, subprocess, sys
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from abas import ABA_EQUIPE, ABA_MAPA
from importar_agenda import importar

SP = os.environ.get("SAIDA", os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(RAIZ, "Centro_Cirurgico_Escala.xlsx")
N_ESP, N_EQ, XC1 = 20, 30, 8
AFC1 = XC1 + N_ESP + 3        # primeira coluna da matriz
AFN_C = XC1 + N_ESP + 2       # auxiliar: quantos X tem a linha
EQ_R1, PAI_R1, PAI_R2 = 2, 10, 25

res, ok = [], 0
def chk(nome, esperado, obtido):
    global ok
    bate = esperado == obtido
    ok += bate
    res.append("%-46s esperado %-24s obtido %-24s %s"
               % (nome, esperado, obtido, "ok" if bate else "DIVERGE"))

def nomes_de(eqp):
    return [eqp.cell(row=EQ_R1 + i, column=2).value for i in range(N_EQ)]

def marcado(eqp, i, j):
    return str(eqp.cell(row=EQ_R1 + i, column=AFC1 + j).value or "").strip().upper() == "X"

def proibidas(eqp):
    n = nomes_de(eqp)
    return {frozenset((n[i], n[j])) for i in range(N_EQ) for j in range(i + 1, N_EQ)
            if n[i] and n[j] and not (marcado(eqp, i, j) and marcado(eqp, j, i))}

def duplas_do_mapa(wb):
    p = wb[ABA_MAPA]
    return [(p.cell(row=r, column=4).value, p.cell(row=r, column=5).value)
            for r in range(PAI_R1, PAI_R2 + 1)]

from datetime import date
CIRS = [dict(sala="Sala %d" % s, data=date(2026, 8, 20), ini="07:00", fim="09:00",
             cirurgiao="Dr. Almeida", procedimento="PROCEDIMENTO FICTICIO", tipo="")
        for s in (1, 2, 3)]

def importa(edita, tag):
    """Copia o arquivo, deixa a edição mexer na matriz, importa e devolve o workbook salvo."""
    alvo = os.path.join(SP, "t10_%s.xlsx" % tag)
    shutil.copy(SRC, alvo)
    wb = openpyxl.load_workbook(alvo)
    edita(wb[ABA_EQUIPE])
    wb.save(alvo)
    saida = importar(None, alvo, os.path.join(SP, "t10_%s_out.xlsx" % tag), cirs=CIRS)
    return openpyxl.load_workbook(saida)

# ---------------------------------------------------------------- a matriz como ela nasce
base = openpyxl.load_workbook(SRC)
eqp0 = base[ABA_EQUIPE]
nomes = nomes_de(eqp0)
usados = [i for i, n in enumerate(nomes) if n]
vazios = [i for i, n in enumerate(nomes) if not n]

chk("a diagonal é a pessoa com ela mesma", {"—"},
    {eqp0.cell(row=EQ_R1 + i, column=AFC1 + i).value for i in range(N_EQ)})
chk("fora da diagonal só há X ou vazio", set(),
    {str(eqp0.cell(row=EQ_R1 + i, column=AFC1 + j).value or "").strip().upper()
     for i in range(N_EQ) for j in range(N_EQ) if i != j} - {"X", ""})
chk("linha ainda sem nome já vem toda marcada", True,
    all(marcado(eqp0, vazios[0], j) for j in range(N_EQ) if j != vazios[0]))
chk("coluna ainda sem nome já vem toda marcada", True,
    all(marcado(eqp0, i, vazios[0]) for i in range(N_EQ) if i != vazios[0]))
chk("o arquivo de demonstração traz duplas sem afinidade", 2, len(proibidas(eqp0)))
chk("a matriz é simétrica no arquivo gerado", [],
    [(nomes[i], nomes[j]) for i in range(N_EQ) for j in range(N_EQ)
     if i != j and marcado(eqp0, i, j) != marcado(eqp0, j, i)])

# ---------------------------------------------------------------- leitura das duplas proibidas
a, b = usados[0], usados[1]
c, d = usados[2], usados[3]

def _apaga_um_lado(eqp):
    eqp.cell(row=EQ_R1 + a, column=AFC1 + b).value = None      # só o lado de cima
wb1 = importa(_apaga_um_lado, "umlado")
chk("apagar o X de um lado proíbe a dupla", True,
    frozenset((nomes[a], nomes[b])) in proibidas(wb1[ABA_EQUIPE]))
chk("apagar de um lado: não ficam juntas no mapa", [],
    [x for x in duplas_do_mapa(wb1) if set(x) == {nomes[a], nomes[b]}])
chk("apagar de um lado: as duas continuam escaladas", True,
    all(n in [x for par in duplas_do_mapa(wb1) for x in par] for n in (nomes[a], nomes[b])))

def _apaga_outro_lado(eqp):
    eqp.cell(row=EQ_R1 + d, column=AFC1 + c).value = None      # só o lado de baixo
wb2 = importa(_apaga_outro_lado, "outrolado")
chk("apagar o X do lado de baixo proíbe a dupla", True,
    frozenset((nomes[c], nomes[d])) in proibidas(wb2[ABA_EQUIPE]))
chk("lado de baixo: não ficam juntas no mapa", [],
    [x for x in duplas_do_mapa(wb2) if set(x) == {nomes[c], nomes[d]}])

def _minusculo(eqp):
    eqp.cell(row=EQ_R1 + a, column=AFC1 + b).value = "x"       # x minúsculo vale como marcado
wb3 = importa(_minusculo, "minusculo")
chk("x minúsculo conta como marcado", False,
    frozenset((nomes[a], nomes[b])) in proibidas(wb3[ABA_EQUIPE]))

# ---------------------------------------------------------------- linha inteira apagada
def _isola(eqp):
    for j in range(N_EQ):
        if j != a:
            eqp.cell(row=EQ_R1 + a, column=AFC1 + j).value = None
wb4 = importa(_isola, "isolada")
mapa4 = duplas_do_mapa(wb4)
comp = [par for par in mapa4 if nomes[a] in par and par[0] and par[1]]
chk("quem está proibida com todas não forma dupla", [], comp)
chk("quem está proibida com todas continua escalada sozinha", True,
    any(nomes[a] in par for par in mapa4))
chk("a isolada fica proibida com todas as cadastradas", len(usados) - 1,
    len([p for p in proibidas(wb4[ABA_EQUIPE]) if nomes[a] in p]))

# ------------------------------------------- a fórmula do Mapa do Dia com o X apagado de um lado
def recalc(path):
    out = os.path.join(SP, "rc10"); shutil.rmtree(out, ignore_errors=True); os.makedirs(out)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx",
                    "--outdir", out, path], capture_output=True)
    return os.path.join(out, os.path.basename(path))

mao = os.path.join(SP, "t10_mao.xlsx")
shutil.copy(SRC, mao)
wbm = openpyxl.load_workbook(mao)
wbm[ABA_EQUIPE].cell(row=EQ_R1 + a, column=AFC1 + b).value = None      # só o lado de cima
pm = wbm[ABA_MAPA]
pm.cell(row=PAI_R1, column=4).value = nomes[a]                          # dupla formada à mão
pm.cell(row=PAI_R1, column=5).value = nomes[b]
wbm.save(mao)
wbr = openpyxl.load_workbook(recalc(mao), data_only=True)
pr = wbr[ABA_MAPA]
chk("X apagado de um lado: a fórmula acusa DUPLA VETADA", "DUPLA VETADA",
    pr.cell(row=PAI_R1, column=6).value)
chk("o alerta nomeia as duas", True,
    str(pr.cell(row=PAI_R1, column=7).value or "").startswith(
        "%s e %s não podem trabalhar juntas" % (nomes[a], nomes[b])))
chk("nenhum erro de fórmula no arquivo editado à mão", [],
    [(w.title, c.coordinate) for w in wbr.worksheets for row in w.iter_rows() for c in row
     if isinstance(c.value, str) and any(e in c.value for e in
        ("#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "Err:"))])

# ------------------------------------------- a verificação do Painel com uma técnica isolada
iso = os.path.join(SP, "t10_isolada_out.xlsx")
wbi = openpyxl.load_workbook(recalc(iso), data_only=True)
pi = wbi[ABA_MAPA]
achou = next((pi.cell(row=r, column=2).value for r in range(1, 90)
              if str(pi.cell(row=r, column=1).value or "").startswith("Técnicos sem afinidade")), None)
chk("verificação: técnico sem afinidade com ninguém", 1, achou)

print("\n".join(res))
print("\nTESTE 10 (afinidade): %d/%d" % (ok, len(res)))
sys.exit(0 if ok == len(res) else 1)
