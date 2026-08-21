# -*- coding: utf-8 -*-
"""Teste 11 — troca da escala de folgas, que é o que muda todo mês.

O que este teste cobra:
  · lançar uma folga DEPOIS de a agenda já ter sido importada não passa despercebido: quem
    ficou de folga e continuou escalado cai em INDISPONÍVEL, com o nome na fila de ações,
    e os indicadores do dia (disponíveis, déficit) reagem na hora;
  · importar a agenda de novo redistribui e ninguém de folga sobra na escala;
  · trocar o primeiro dia do período na Configuração move o calendário inteiro.
"""
import os, shutil, subprocess, sys
from datetime import date, datetime, timedelta
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from abas import ABA_CFG, ABA_EQUIPE, ABA_FOLGAS, ABA_MAPA, ABA_PAINEL
from importar_agenda import importar

SP = os.environ.get("SAIDA", os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(RAIZ, "Centro_Cirurgico_Escala.xlsx")
PAI_R1, PAI_R2, AUS_R1, AUS_R2, CAL_C1 = 10, 25, 2, 151, 10

res, ok = [], 0
def chk(nome, esperado, obtido):
    global ok
    bate = esperado == obtido
    ok += bate
    res.append("%-46s esperado %-26s obtido %-26s %s"
               % (nome, esperado, obtido, "ok" if bate else "DIVERGE"))

def recalc(path, tag):
    out = os.path.join(SP, "rc11_" + tag)
    shutil.rmtree(out, ignore_errors=True); os.makedirs(out)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx",
                    "--outdir", out, path], capture_output=True)
    return openpyxl.load_workbook(os.path.join(out, os.path.basename(path)), data_only=True)

def duplas(wb):
    p = wb[ABA_MAPA]
    return [(p.cell(row=r, column=4).value, p.cell(row=r, column=5).value)
            for r in range(PAI_R1, PAI_R2 + 1)]

def fila(wb):
    d = wb[ABA_PAINEL]
    return [str(d.cell(row=r, column=3).value) for r in range(20, 60)
            if d.cell(row=r, column=3).value and "·" in str(d.cell(row=r, column=3).value)]

def linha_livre(aus):
    return next(r for r in range(AUS_R1, AUS_R2 + 1) if aus.cell(row=r, column=1).value is None)

DIA = openpyxl.load_workbook(SRC)[ABA_MAPA]["A6"].value
DIA = DIA.date() if isinstance(DIA, datetime) else DIA
CIRS = [dict(sala="Sala %d" % s, data=DIA, ini="07:00", fim="09:00",
             cirurgiao="Dr. Almeida", procedimento="PROCEDIMENTO FICTICIO", tipo="")
        for s in (1, 2, 3, 4)]

# ---------------------------------------------------- estado de partida: agenda já importada
base = os.path.join(SP, "t11_base.xlsx")
shutil.copy(SRC, base)
importado = importar(None, base, os.path.join(SP, "t11_importado.xlsx"), cirs=CIRS)
antes = recalc(importado, "antes")
p0 = antes[ABA_MAPA]
disp0, def0 = p0["E6"].value, p0["G6"].value
alvo = next(t for t, _ in duplas(antes) if t)
chk("ninguém em INDISPONÍVEL antes da folga", 0,
    sum(1 for r in range(PAI_R1, PAI_R2 + 1)
        if p0.cell(row=r, column=6).value == "INDISPONÍVEL"))

# ---------------------------------------------------- 1) folga lançada DEPOIS da importação
comfolga = os.path.join(SP, "t11_comfolga.xlsx")
shutil.copy(importado, comfolga)
wb = openpyxl.load_workbook(comfolga)
aus = wb[ABA_FOLGAS]
r = linha_livre(aus)
aus.cell(row=r, column=1).value = alvo
aus.cell(row=r, column=2).value = "Folga"
aus.cell(row=r, column=3).value = DIA
aus.cell(row=r, column=4).value = DIA
wb.save(comfolga)
depois = recalc(comfolga, "depois")
pd_ = depois[ABA_MAPA]

chk("a folga aparece na aba Equipe", "Folga",
    next(pd_.cell(row=rr, column=2).value for rr in range(1, 200)
         if pd_.cell(row=rr, column=1).value == alvo and rr > 26))
chk("disponíveis caem em 1", disp0 - 1, pd_["E6"].value)
chk("déficit sobe em 1", def0 + 1, pd_["G6"].value)
chk("o ambiente dela vira INDISPONÍVEL", "INDISPONÍVEL",
    next(pd_.cell(row=rr, column=6).value for rr in range(PAI_R1, PAI_R2 + 1)
         if alvo in (pd_.cell(row=rr, column=4).value, pd_.cell(row=rr, column=5).value)))
chk("o alerta diz o nome e o motivo", True,
    any((alvo + " não está disponível hoje (Folga)") in a for a in
        [str(pd_.cell(row=rr, column=7).value) for rr in range(PAI_R1, PAI_R2 + 1)]))
chk("a fila de ações do Painel avisa", True,
    any(alvo in t and "não está disponível" in t for t in fila(depois)))
chk("a verificação conta o ambiente", 1,
    next(pd_.cell(row=rr, column=2).value for rr in range(1, 90)
         if str(pd_.cell(row=rr, column=1).value or "").startswith("Ambientes com alguém de folga")))

# ---------------------------------------------------- 2) reimportar redistribui
redistribuido = importar(None, comfolga, os.path.join(SP, "t11_redistribuido.xlsx"), cirs=CIRS)
novo = recalc(redistribuido, "novo")
escalados = [n for par in duplas(novo) for n in par if n]
chk("depois de reimportar, ela sai da escala", False, alvo in escalados)
pn = novo[ABA_MAPA]
chk("ninguém de folga sobra escalado", 0,
    sum(1 for rr in range(PAI_R1, PAI_R2 + 1)
        if pn.cell(row=rr, column=6).value == "INDISPONÍVEL"))
chk("a verificação zera", 0,
    next(pn.cell(row=rr, column=2).value for rr in range(1, 90)
         if str(pn.cell(row=rr, column=1).value or "").startswith("Ambientes com alguém de folga")))

# ---------------------------------------------------- 3) trocar o período da escala
novo_per = os.path.join(SP, "t11_periodo.xlsx")
shutil.copy(SRC, novo_per)
wb = openpyxl.load_workbook(novo_per)
p_ini = wb[ABA_CFG]["B8"].value
p_ini = p_ini.date() if isinstance(p_ini, datetime) else p_ini
alvo_per = date(p_ini.year, p_ini.month + 1, p_ini.day) if p_ini.month < 12 else \
           date(p_ini.year + 1, 1, p_ini.day)
wb[ABA_CFG]["B8"] = alvo_per
wb.save(novo_per)
per = recalc(novo_per, "periodo")
ca = per[ABA_FOLGAS]
datas = [ca.cell(row=1, column=CAL_C1 + d).value for d in range(31)]
datas = [x.date() if isinstance(x, datetime) else x for x in datas]
chk("o calendário começa no novo período", alvo_per, datas[0])
chk("o calendário anda 31 dias seguidos", 31,
    sum(1 for d in range(31) if datas[d] == alvo_per + timedelta(days=d)))

print("\n".join(res))
print("\nTESTE 11 (troca da escala de folgas): %d/%d" % (ok, len(res)))
sys.exit(0 if ok == len(res) else 1)
