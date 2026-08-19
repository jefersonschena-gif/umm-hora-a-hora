import openpyxl, sys
f=sys.argv[1]
wb=openpyxl.load_workbook(f, data_only=True)
ERR=("#REF!","#VALUE!","#DIV/0!","#N/A","#NAME?","#NUM!","#NULL!","Err:")
bad=[]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and any(e in v for e in ERR):
                bad.append((ws.title,c.coordinate,v))
print("ERROS DE FÓRMULA:", len(bad))
for b in bad[:25]: print("  ",b)
esc=wb["Escala_Duplas"]
print("\n%-7s %-6s %3s %5s %6s %-22s %-22s %5s %5s %5s %-12s %s"%("SALA","TURNO","N","MIN","OCUP","CIRC","INSTR","AC","AI","DUP","CLASS","ALERTAS"))
for r in range(2,38):
    b=esc.cell(row=r,column=2).value
    if not b: continue
    g=esc.cell(row=r,column=7).value
    vals=[esc.cell(row=r,column=c).value for c in (2,3,4,5,8,10,12,13,14,15,20)]
    print("%-7s %-6s %3s %5s %6s %-22s %-22s %5s %5s %5s %-12s %s"%(
        vals[0],vals[1],vals[2],vals[3], ("%.0f%%"%(g*100)) if isinstance(g,(int,float)) else "-",
        str(vals[4])[:22],str(vals[5])[:22],
        ("%.2f"%vals[6]) if isinstance(vals[6],(int,float)) else "-",
        ("%.2f"%vals[7]) if isinstance(vals[7],(int,float)) else "-",
        ("%.2f"%vals[8]) if isinstance(vals[8],(int,float)) else "-",
        vals[9], vals[10]))
pai=wb["Painel"]
print("\nVERIFICAÇÕES:")
for r in range(53,70):
    a=pai.cell(row=r,column=1).value; c=pai.cell(row=r,column=3).value
    if a: print("  %-58s %s"%(a,c))
