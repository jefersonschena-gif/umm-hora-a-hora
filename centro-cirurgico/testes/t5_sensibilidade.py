# -*- coding: utf-8 -*-
"""Teste 5 — sensibilidade: altera premissas e confere a propagação até os indicadores."""
import openpyxl, subprocess, os, shutil
SRC=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                 "Centro_Cirurgico_Escala_Duplas_v1.xlsx")
SP=os.environ.get("SAIDA", os.path.dirname(os.path.abspath(__file__)))

def recalc(path):
    out=os.path.join(SP,"rc5"); shutil.rmtree(out,ignore_errors=True); os.makedirs(out)
    subprocess.run(["soffice","--headless","--norestore","--convert-to","xlsx","--outdir",out,path],
                   capture_output=True)
    return openpyxl.load_workbook(os.path.join(out,os.path.basename(path)),data_only=True)

def cenario(tag, muda):
    wb=openpyxl.load_workbook(SRC)
    for cel,v in muda.items(): wb["Parâmetros"][cel]=v
    f=os.path.join(SP,"t5_%s.xlsx"%tag); wb.save(f)
    return recalc(f)

base=recalc(SRC)
def snap(wb):
    e=wb["Escala_Duplas"]; p=wb["Painel"]
    return dict(
        afin={(e.cell(row=r,column=2).value,e.cell(row=r,column=3).value):e.cell(row=r,column=14).value
              for r in range(2,38) if e.cell(row=r,column=2).value},
        instr={(e.cell(row=r,column=2).value,e.cell(row=r,column=3).value):e.cell(row=r,column=13).value
              for r in range(2,38) if e.cell(row=r,column=2).value},
        ocup={(e.cell(row=r,column=2).value,e.cell(row=r,column=3).value):e.cell(row=r,column=7).value
              for r in range(2,38) if e.cell(row=r,column=2).value},
        classe=[e.cell(row=r,column=15).value for r in range(2,38)],
        kpi=[p.cell(row=6,column=c).value for c in range(2,9)],
        alertas=sum(1 for r in range(2,38) if e.cell(row=r,column=20).value not in (None,"","OK")))
b=snap(base)
print("BASE  KPI:",["%.4g"%v if isinstance(v,float) else v for v in b["kpi"]],
      "| ADEQUADA:",b["classe"].count("ADEQUADA"),"ATENÇÃO:",b["classe"].count("ATENÇÃO"),
      "CRÍTICA:",b["classe"].count("CRÍTICA"),"| alertas:",b["alertas"])
falhas=[]

# 1) peso do instrumentador = 100% -> afinidade da dupla == afinidade do instrumentador
s=snap(cenario("pesos",{"B17":1.0,"B18":0.0}))
d=[k for k in s["afin"] if s["afin"][k] not in (None,"") and abs(s["afin"][k]-s["instr"][k])>1e-9]
print("1) pesos 1,00/0,00 -> afinidade da dupla = do instrumentador:", "OK" if not d else "FALHOU %s"%d[:3])
falhas += d

# 2) meta de afinidade 2,95 -> duplas com afinidade < 2,95 saem de ADEQUADA
for meta,minimo in ((2.95,2.0),(3.50,2.90)):
    s=snap(cenario("meta%s"%meta,{"B20":meta,"B21":minimo}))
    vals=[v for v in b["afin"].values() if isinstance(v,(int,float))]
    esp_adeq=sum(1 for v in vals if v>=meta)
    esp_aten=sum(1 for v in vals if minimo<=v<meta)
    esp_crit=sum(1 for v in vals if v<minimo)
    obt=(s["classe"].count("ADEQUADA"),s["classe"].count("ATENÇÃO"),s["classe"].count("CRÍTICA"))
    ok = obt==(esp_adeq,esp_aten,esp_crit)
    print("2) meta=%.2f min=%.2f -> ADEQUADA/ATENÇÃO/CRÍTICA %s (esperado %s)"%(meta,minimo,obt,(esp_adeq,esp_aten,esp_crit)),
          "OK" if ok else "FALHOU")
    if not ok: falhas.append("meta%s"%meta)

# 3) setup 30->15 min -> ocupação cai; ocupação nova = (ocup*minTurno - 15*nCir)/minTurno
s=snap(cenario("setup",{"B26":15}))
e=base["Escala_Duplas"]
ruim=[]
for r in range(2,38):
    sala=e.cell(row=r,column=2).value
    if not sala: continue
    t=e.cell(row=r,column=3).value; n=e.cell(row=r,column=4).value
    mt=e.cell(row=r,column=6).value
    if not isinstance(b["ocup"][(sala,t)],float): continue
    esperado=(b["ocup"][(sala,t)]*mt-15*n)/mt
    if abs(s["ocup"][(sala,t)]-esperado)>1e-9: ruim.append((sala,t,esperado,s["ocup"][(sala,t)]))
print("3) Setup_Min 30->15 -> ocupação recalculada:", "OK" if not ruim else "FALHOU %s"%ruim[:2])
falhas += ruim

# 4) participação mínima crítica 20%->60% -> menos alertas de inaptidão
s=snap(cenario("part",{"B23":0.60}))
print("4) Part_Min_Critica 20%%->60%% -> alertas: %d (base %d)"%(s["alertas"],b["alertas"]),
      "OK" if s["alertas"]<=b["alertas"] else "FALHOU")

# 5) restauração: arquivo original intacto
s=snap(recalc(SRC))
igual = s["afin"]==b["afin"] and s["kpi"]==b["kpi"]
print("5) arquivo original inalterado após os testes:", "OK" if igual else "FALHOU")
print("\nRESULTADO TESTE 5:", "APROVADO" if not falhas and igual else "REPROVADO")
